# excel.py
import openpyxl
from openpyxl.chart import BarChart, Reference
from database import (
    get_expenses_for_month, get_savings_goals, get_debts, 
    get_all_category_budgets, get_all_incomes_month
)
from datetime import datetime

def export_excel(db_path, path_xlsx):
    wb = openpyxl.Workbook()
    
    # --- Hoja 1: Gastos ---
    ws_gastos = wb.active
    ws_gastos.title = "Gastos"
    ws_gastos.append(["Fecha", "Descripción", "Categoría", "Monto"])
    mes_actual = datetime.now().strftime('%Y-%m')
    gastos = get_expenses_for_month(db_path, mes_actual)
    for g in gastos:
        # g = (id, created_at, desc, amount, cat, currency)
        ws_gastos.append([g[1][:10], g[2], g[4], g[3]])

    # --- Hoja 2: Metas de Ahorro ---
    ws_metas = wb.create_sheet("Metas de Ahorro")
    ws_metas.append(["Nombre", "Meta", "Ahorrado"])
    for m in get_savings_goals(db_path):
        ws_metas.append([m[1], m[2], m[3]])

    # --- Hoja 3: Deudas ---
    ws_deudas = wb.create_sheet("Deudas")
    ws_deudas.append(["Persona", "Monto", "Tipo", "Estado"])
    for d in get_debts(db_path, "pending"):
        ws_deudas.append([d[1], d[2], d[3], d[4]])

    # --- Hoja 4: Presupuestos ---
    ws_presupuestos = wb.create_sheet("Presupuestos")
    ws_presupuestos.append(["Categoría", "Límite"])
    for p in get_all_category_budgets(db_path):
        ws_presupuestos.append([p[0], p[1]])

    # --- Hoja 5: Ingresos Extra ---
    ws_ingresos = wb.create_sheet("Ingresos Extra")
    ws_ingresos.append(["Descripción", "Monto", "Fecha"])
    for i in get_all_incomes_month(db_path, mes_actual):
        ws_ingresos.append([i[1], i[2], i[3]])

    # --- Gráfico de Gastos por Categoría ---
    ws_grafico = wb.create_sheet("Gráfico Gastos")
    ws_grafico.append(["Categoría", "Monto"])
    # Agrupar gastos por categoría
    cat_totals = {}
    for g in gastos:
        cat = g[4]
        monto = g[3]
        cat_totals[cat] = cat_totals.get(cat, 0) + monto
    
    for cat, total in cat_totals.items():
        ws_grafico.append([cat, total])

    chart = BarChart()
    chart.title = "Gastos por Categoría"
    chart.y_axis.title = 'Monto'
    chart.x_axis.title = 'Categoría'
    
    data = Reference(ws_grafico, min_col=2, min_row=1, max_row=len(cat_totals)+1)
    cats = Reference(ws_grafico, min_col=1, min_row=2, max_row=len(cat_totals)+1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    ws_grafico.add_chart(chart, "D2")

    wb.save(path_xlsx)
